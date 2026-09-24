from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

prs = Presentation()

# Layouts
TITLE_SLIDE = 0
BULLET_SLIDE = 1
SECTION_HEADER = 2

# Helper to apply styling
def style_title(shape):
    for run in shape.text_frame.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 112, 192)  # Blue

# Slide 1: Title
slide_layout = prs.slide_layouts[TITLE_SLIDE]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "HealthHadoop AI"
subtitle.text = "Next-Generation Big Data Healthcare Analytics\nCapstone Project Presentation"
style_title(title)

# Slide 2: The Problem
slide_layout = prs.slide_layouts[BULLET_SLIDE]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "The Challenge in Healthcare"
tf = body.text_frame
tf.text = "Hospitals struggle with fragmented and overwhelming data:"
p = tf.add_paragraph()
p.text = "Reactive Patient Care: Interventions happen after vitals drop, rather than predicting deterioration."
p.level = 1
p = tf.add_paragraph()
p.text = "Data Silos: Financial records, patient demographics, and clinical outcomes are disconnected."
p.level = 1
p = tf.add_paragraph()
p.text = "Slow Reporting: Batch processing takes too long to generate actionable insights for doctors."
p.level = 1

# Slide 3: The Solution
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "The Solution: HealthHadoop AI"
tf = body.text_frame
tf.text = "A unified, AI-powered analytics dashboard that bridges the gap between historical data and real-time monitoring."
p = tf.add_paragraph()
p.text = "Batch Analytics: Instant ETL processing for long-term trends and cost analysis."
p.level = 1
p = tf.add_paragraph()
p.text = "Real-Time Streaming: Live ICU patient vitals monitoring with anomaly detection."
p.level = 1
p = tf.add_paragraph()
p.text = "AI Data Studio: Chat-based LLM integration to query complex data effortlessly."
p.level = 1

# Slide 4: Key Features & UI
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "Key Features & Dashboard"
tf = body.text_frame
tf.text = "Built with a premium 'Aurora Glass' UI for maximum clinical visibility:"
p = tf.add_paragraph()
p.text = "Custom CSV Uploads: In-memory Pandas processing scales dynamically to user data."
p.level = 1
p = tf.add_paragraph()
p.text = "Bento-Grid Visualizations: Interactive charts for Regional Burden, Disease Trends, and Readmission Rates."
p.level = 1
p = tf.add_paragraph()
p.text = "Predictive Modeling: Built-in Random Forest inference to predict 30-day patient readmission risk."
p.level = 1

# Slide 5: System Architecture
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "System Architecture"
tf = body.text_frame
tf.text = "Modern, resilient, and optimized for cloud deployment:"
p = tf.add_paragraph()
p.text = "Frontend: React, Vite, Tailwind CSS, Framer Motion, and Recharts."
p.level = 1
p = tf.add_paragraph()
p.text = "Backend API: FastAPI (Python) with JWT Authentication."
p.level = 1
p = tf.add_paragraph()
p.text = "Data Pipeline: High-speed Pandas in-memory processing (replacing heavy disk I/O)."
p.level = 1
p = tf.add_paragraph()
p.text = "Streaming: Python AsyncIO WebSockets simulating high-frequency Kafka telemetry."
p.level = 1

# Slide 6: Business Impact
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
body = slide.placeholders[1]
title.text = "Business Impact"
tf = body.text_frame
tf.text = "Transforming data into better patient outcomes:"
p = tf.add_paragraph()
p.text = "Reduced Readmissions: High-risk patient identification allows for preventative care."
p.level = 1
p = tf.add_paragraph()
p.text = "Cost Optimization: Clear visibility into average treatment costs across disease categories."
p.level = 1
p = tf.add_paragraph()
p.text = "Faster Triage: Real-time anomaly detection alerts nurses instantly to critical drops in oxygen or spikes in heart rate."
p.level = 1

# Slide 7: Conclusion
slide_layout = prs.slide_layouts[SECTION_HEADER]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Thank You!"
subtitle.text = "Live Demo & Q&A"
style_title(title)

prs.save("HealthHadoop_AI_Presentation.pptx")
print("Presentation generated successfully!")
