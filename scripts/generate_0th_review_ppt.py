import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

prs = Presentation()

# Layouts
TITLE_SLIDE = 0
BULLET_SLIDE = 1

def add_slide(title_text, bullets):
    slide_layout = prs.slide_layouts[BULLET_SLIDE]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    
    # Style title
    for run in title.text_frame.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 112, 192)

    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()  # clear default text
    
    for bullet in bullets:
        p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(20)

# 1. Title Slide
slide_layout = prs.slide_layouts[TITLE_SLIDE]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "HealthHadoop AI: Big Data Healthcare Analytics"
for run in title.text_frame.paragraphs[0].runs:
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 112, 192)

subtitle.text = (
    "0th Review - BDE Project\n\n"
    "Team Members: [Your Names]\n"
    "Register Numbers: [Your Register Numbers]\n"
    "Department / College: [Your Department & College]"
)

# 2. Abstract
add_slide("Abstract", [
    "Develops a scalable Big Data healthcare analytics platform combining real-time streaming and batch processing.",
    "Problem: Healthcare data is fragmented, leading to reactive patient care and delayed insights.",
    "Solution: A Lambda-architecture-inspired dashboard using in-memory data processing and simulated streaming for instant clinical insights.",
    "Technologies: React, FastAPI, Pandas, Python WebSockets, Grok AI.",
    "Expected Outcome: A responsive, predictive dashboard that forecasts 30-day readmissions and tracks disease burden in real-time."
])

# 3. Introduction
add_slide("Introduction", [
    "Background: The healthcare industry generates massive amounts of data daily, yet much of it remains siloed and underutilized.",
    "Importance: Rapid analysis of clinical and financial data is critical for saving lives and optimizing hospital resources.",
    "Current situation: Many hospitals rely on legacy batch-processing systems that delay critical alerts, such as sudden drops in patient vitals.",
    "Motivation: To bridge the gap between long-term historical data analysis and real-time patient monitoring in a unified, intuitive platform."
])

# 4. Problem Statement
add_slide("Problem Statement", [
    "Existing Problem: Inability to dynamically analyze massive healthcare datasets in real-time to predict patient readmissions and monitor ICU vitals.",
    "Limitations of current approaches: Existing systems are either exclusively batch-oriented (slow) or lack predictive AI integration, making them purely reactive.",
    "Who is affected: Medical staff (doctors/nurses) who lack immediate actionable insights, and patients who suffer from delayed preventative care."
])

# 5. Objectives
add_slide("Objectives", [
    "Main Objective: To design and implement a comprehensive Big Data healthcare analytics platform for real-time monitoring and predictive insights.",
    "Specific Objective 1: Develop an in-memory batch processing pipeline for rapid ETL on large healthcare datasets.",
    "Specific Objective 2: Implement a real-time streaming layer to simulate and monitor ICU patient vitals.",
    "Specific Objective 3: Integrate a Machine Learning model to predict 30-day patient readmission risks.",
    "Specific Objective 4: Incorporate a conversational AI assistant (Grok AI) to allow intuitive querying of complex medical data."
])

# 6. Existing System / Existing Work
add_slide("Existing System / Existing Work", [
    "Current methods: Traditional RDBMS (Relational Databases) and legacy dashboards that run overnight batch jobs.",
    "Existing technologies: Standard SQL reporting and basic static BI tools (e.g., Tableau, PowerBI) without real-time streaming integration.",
    "Limitations: High latency for actionable insights, inability to handle high-velocity streaming telemetry data, and lack of built-in predictive ML models.",
    "Research gap: A unified architectural framework that seamlessly integrates high-speed streaming (telemetry) and batch processing (historical records) with conversational AI."
])

# 7. Proposed System
add_slide("Proposed System", [
    "Proposed Solution: 'HealthHadoop AI', a modern web-based platform utilizing FastAPI and React for seamless data orchestration.",
    "How it improves: Replaces heavy disk I/O with high-speed in-memory Pandas processing and introduces asynchronous WebSockets for zero-latency vitals tracking.",
    "Key Feature 1: Bento-Grid Interactive Dashboard for regional and demographic metrics.",
    "Key Feature 2: Live Patient Streaming Layer with anomaly detection.",
    "Key Feature 3: Predictive Readmission Risk Calculator (ML).",
    "Key Feature 4: AI Data Studio for NLP-based data exploration."
])

prs.save("BDE_0th_Review_Presentation.pptx")
print("Presentation generated successfully!")
